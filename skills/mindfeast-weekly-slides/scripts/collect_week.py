#!/usr/bin/env python3
"""Collect one student's logged lessons for a seven-day period."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import openpyxl
import yaml

DATE_HEADERS = {"date"}
START_HEADERS = {"start time", "start"}
END_HEADERS = {"end time", "end"}
DESCRIPTION_HEADERS = {"description", "lesson", "task"}
TEACHER_HEADERS = {"teacher"}
NOTES_HEADERS = {"notes", "note"}


def normalize_header(value: Any) -> str:
    """Return a lowercase, whitespace-normalized column header."""
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def parse_date(value: Any) -> date | None:
    """Parse spreadsheet date values into a date, returning None on blank or unknown values."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%m/%d/%Y", "%-m/%-d/%Y", "%Y-%m-%d", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def cell_text(value: Any) -> str:
    """Convert a spreadsheet cell value into stable display text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%-m/%-d/%Y")
    if isinstance(value, date):
        return value.strftime("%-m/%-d/%Y")
    return str(value).strip()


def time_sort_key(value: str) -> tuple[int, int, str]:
    """Return a sortable key for common spreadsheet time formats."""
    text = value.strip()
    for fmt in ("%I:%M %p", "%I %p", "%H:%M"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.hour, parsed.minute, text
        except ValueError:
            continue
    return 99, 99, text


def find_column(headers: list[str], accepted: set[str]) -> int | None:
    """Find the first header whose normalized name is in the accepted set."""
    for index, header in enumerate(headers):
        if header in accepted:
            return index
    return None


def rolling_week_start(today: date, days: int = 7) -> date:
    """Return the first day of an inclusive rolling window ending today."""
    return today - timedelta(days=days - 1)


def resolve_student(students: dict[str, Any], query: str) -> tuple[str, dict[str, Any]]:
    """Resolve a student by slug, display name, or alias."""
    needle = query.strip().lower()
    for slug, data in students.items():
        candidates = {slug.lower(), str(data.get("display_name", "")).lower()}
        candidates.update(str(alias).lower() for alias in data.get("aliases", []) or [])
        if needle in candidates:
            return slug, data
    raise SystemExit(f"student not found: {query}")


def collect_rows(
    workbook_path: Path, subjects: list[str], start: date, end: date
) -> list[dict[str, str]]:
    """Collect logged lesson rows from subject sheets within the inclusive date range."""
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    rows: list[dict[str, str]] = []
    subject_names = subjects or wb.sheetnames

    for subject in subject_names:
        if subject not in wb.sheetnames:
            continue
        ws = wb[subject]
        iterator = ws.iter_rows(values_only=True)
        try:
            header_row = next(iterator)
        except StopIteration:
            continue

        headers = [normalize_header(value) for value in header_row]
        date_col = find_column(headers, DATE_HEADERS)
        start_col = find_column(headers, START_HEADERS)
        end_col = find_column(headers, END_HEADERS)
        desc_col = find_column(headers, DESCRIPTION_HEADERS)
        teacher_col = find_column(headers, TEACHER_HEADERS)
        notes_col = find_column(headers, NOTES_HEADERS)

        if date_col is None or desc_col is None:
            continue

        for row in iterator:
            row_date = parse_date(row[date_col] if date_col < len(row) else None)
            if row_date is None or row_date < start or row_date > end:
                continue
            description = cell_text(row[desc_col] if desc_col < len(row) else "")
            notes = cell_text(
                row[notes_col] if notes_col is not None and notes_col < len(row) else ""
            )
            if not description and not notes:
                continue
            rows.append(
                {
                    "subject": subject,
                    "date": row_date.isoformat(),
                    "start_time": cell_text(
                        row[start_col]
                        if start_col is not None and start_col < len(row)
                        else ""
                    ),
                    "end_time": cell_text(
                        row[end_col]
                        if end_col is not None and end_col < len(row)
                        else ""
                    ),
                    "description": description,
                    "teacher": cell_text(
                        row[teacher_col]
                        if teacher_col is not None and teacher_col < len(row)
                        else ""
                    ),
                    "notes": notes,
                }
            )

    rows.sort(
        key=lambda item: (
            item["date"],
            time_sort_key(item["start_time"]),
            item["subject"],
            item["description"],
        )
    )
    return rows


def main() -> int:
    """Run the CLI and print weekly lesson rows as JSON."""
    parser = argparse.ArgumentParser()
    parser.add_argument("--students-yaml", default="students.yaml", type=Path)
    parser.add_argument("--student", required=True)
    parser.add_argument(
        "--week-start",
        help="Start date in YYYY-MM-DD format; defaults to the last seven days",
    )
    parser.add_argument("--days", type=int, default=7)
    args = parser.parse_args()
    if args.days < 1:
        parser.error("--days must be at least 1")

    registry = yaml.safe_load(args.students_yaml.read_text(encoding="utf-8")) or {}
    slug, student = resolve_student(registry.get("students", {}), args.student)

    today = date.today()
    if args.week_start:
        start = datetime.strptime(args.week_start, "%Y-%m-%d").date()
        end = min(start + timedelta(days=args.days - 1), today)
    else:
        end = today
        start = rolling_week_start(today, args.days)

    spreadsheet_value = student.get("time_tracking_spreadsheet")
    if not spreadsheet_value:
        raise SystemExit(f"time_tracking_spreadsheet missing for {slug}")
    spreadsheet = Path(str(spreadsheet_value)).expanduser()
    if not spreadsheet.is_file():
        raise SystemExit(f"time tracking spreadsheet not found: {spreadsheet}")

    rows = collect_rows(spreadsheet, student.get("subjects", []) or [], start, end)
    result = {
        "student_slug": slug,
        "display_name": student.get("display_name", slug),
        "grade": student.get("grade"),
        "week_start": start.isoformat(),
        "week_end": end.isoformat(),
        "spreadsheet": str(spreadsheet),
        "row_count": len(rows),
        "rows": rows,
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
