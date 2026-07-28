from __future__ import annotations

import argparse
import importlib.util
import re
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml


def load_module():
    repo = Path(__file__).resolve().parents[1]
    scripts_dir = repo / "scripts"
    sys.path.insert(0, str(scripts_dir))
    try:
        spec = importlib.util.spec_from_file_location(
            "hsd_project", scripts_dir / "hsd_project.py"
        )
        assert spec is not None
        assert spec.loader is not None
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module
    finally:
        sys.path.remove(str(scripts_dir))


SCIENCE_HEADERS = ("Date", "Start Time", "End Time", "Description", "Teacher",
                   "Materials", "ISBN", "Notes")
MATH_HEADERS = ("Date", "Start Time", "End Time", "Description", "Teacher",
                "Materials", "ISBN")


def build_workbook(path: Path, extra_rows: dict[str, list[tuple]] | None = None) -> None:
    wb = openpyxl.Workbook()
    science = wb.active
    science.title = "Science"
    science.append(SCIENCE_HEADERS)
    math = wb.create_sheet("Math")
    math.append(MATH_HEADERS)
    for sheet_name, rows in (extra_rows or {}).items():
        for row in rows:
            wb[sheet_name].append(row)
    wb.save(path)


def build_env(
    tmp_path: Path,
    hsd: bool = True,
    time: bool = True,
    lesson_log: bool = True,
) -> dict[str, str]:
    workbook = tmp_path / "Time-3.xlsx"
    build_workbook(workbook)

    registry = {
        "students": {
            "eliana": {
                "display_name": "Eliana",
                "grade": 3,
                "aliases": ["eliana"],
                "subjects": ["Science", "Math"],
                "time_tracking_spreadsheet": str(workbook),
            }
        }
    }
    registry_path = tmp_path / "students.yaml"
    registry_path.write_text(yaml.safe_dump(registry), encoding="utf-8")

    runtime_path = tmp_path / "runtime.yaml"
    runtime_path.write_text(
        yaml.safe_dump(
            {
                "logging": {
                    "lesson_log": lesson_log,
                    "time": time,
                    "hsd": hsd,
                }
            }
        ),
        encoding="utf-8",
    )

    return {
        "root": str(tmp_path / "lesson-logs"),
        "registry": str(registry_path),
        "runtime": str(runtime_path),
        "workbook": str(workbook),
    }


def write_log(env: dict[str, str], *, date: str, subject: str, slug: str,
              start: str | None = "9:15 AM", end: str | None = "10:05 AM",
              teacher: str = "Brad", grade_dir: str = "grade-3") -> Path:
    year, month, day = date.split("-")
    subject_slug = re.sub(r"[^a-z0-9]+", "-", subject.lower()).strip("-")
    session = (
        Path(env["root"])
        / "eliana"
        / grade_dir
        / year
        / month
        / day
        / f"{subject_slug}-{slug}"
    )
    (session / "sources").mkdir(parents=True)
    frontmatter = {
        "schema": "lesson-log/v1",
        "student": "eliana",
        "grade": 3,
        "date": date,
        "subject": subject,
        "teacher": teacher,
    }
    if start:
        frontmatter["start_time"] = start
    if end:
        frontmatter["end_time"] = end
    body = yaml.safe_dump(frontmatter, sort_keys=False).strip()
    (session / "log.md").write_text(f"---\n{body}\n---\n\n# {subject}\n", encoding="utf-8")
    return session


def gaps_args(env: dict[str, str], **kwargs) -> argparse.Namespace:
    base = {
        "root": env["root"],
        "registry": env["registry"],
        "runtime": env["runtime"],
        "student": "eliana",
        "date_from": None,
        "date_to": None,
    }
    base.update(kwargs)
    return argparse.Namespace(**base)


def append_args(env: dict[str, str], log: Path, **kwargs) -> argparse.Namespace:
    base = {
        "root": env["root"],
        "registry": env["registry"],
        "runtime": env["runtime"],
        "student": "eliana",
        "log": str(log),
        "description": "Lesson 12: Mapping a Plot",
        "notes": None,
    }
    base.update(kwargs)
    return argparse.Namespace(**base)


def test_gaps_lists_a_log_with_no_row(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    write_log(env, date="2026-07-24", subject="Science", slug="science-plot")

    result = mod.cmd_gaps(gaps_args(env))

    assert result["logs_examined"] == 1
    assert len(result["pending"]) == 1
    assert result["pending"][0]["subject"] == "Science"
    assert result["pending"][0]["has_notes_column"] is True


def test_gaps_reports_no_notes_column_for_sheets_without_one(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    write_log(env, date="2026-07-24", subject="Math", slug="math-45")

    result = mod.cmd_gaps(gaps_args(env))
    assert result["pending"][0]["has_notes_column"] is False


def test_gaps_skips_sessions_that_already_have_a_row(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    build_workbook(
        Path(env["workbook"]),
        {"Science": [("7/24/2026", "9:15 AM", "10:05 AM", "Lesson 12", "Brad")]},
    )
    write_log(env, date="2026-07-24", subject="Science", slug="science-plot")

    result = mod.cmd_gaps(gaps_args(env))
    assert result["pending"] == []
    assert result["already_logged"] == 1


def test_gaps_separates_logs_that_have_no_times_yet(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    write_log(env, date="2026-07-24", subject="Science", slug="plot", start=None, end=None)

    result = mod.cmd_gaps(gaps_args(env))
    assert result["pending"] == []
    assert len(result["waiting_on_times"]) == 1


def test_gaps_waits_when_only_the_end_time_is_missing(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    write_log(
        env,
        date="2026-07-24",
        subject="Science",
        slug="plot",
        start="9:15 AM",
        end=None,
    )

    result = mod.cmd_gaps(gaps_args(env))
    assert result["pending"] == []
    assert len(result["waiting_on_times"]) == 1


def test_gaps_honours_a_date_range(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    write_log(env, date="2026-07-20", subject="Science", slug="early")
    write_log(env, date="2026-07-24", subject="Science", slug="late")

    result = mod.cmd_gaps(gaps_args(env, date_from="2026-07-22"))
    assert [p["date"] for p in result["pending"]] == ["2026-07-24"]


def test_gaps_ignores_logs_from_a_previous_grade(tmp_path):
    """A grade-2 log belongs to a grade-2 workbook, not this one."""
    mod = load_module()
    env = build_env(tmp_path)
    write_log(env, date="2026-07-24", subject="Science", slug="old", grade_dir="grade-2")

    result = mod.cmd_gaps(gaps_args(env))
    assert result["logs_examined"] == 0


def test_gaps_never_reports_spreadsheet_rows_that_have_no_log(tmp_path):
    """Pre-log history must be invisible to the projection."""
    mod = load_module()
    env = build_env(tmp_path)
    build_workbook(
        Path(env["workbook"]),
        {"Science": [("3/2/2026", "9:00 AM", "9:45 AM", "Old row", "Brad")]},
    )

    result = mod.cmd_gaps(gaps_args(env))
    assert result["pending"] == []
    assert result["already_logged"] == 0


def test_gaps_flags_a_subject_with_no_matching_sheet(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    write_log(env, date="2026-07-24", subject="Astronomy", slug="stars")

    result = mod.cmd_gaps(gaps_args(env))
    assert len(result["unknown_subject"]) == 1


def test_append_writes_the_row_with_teacher_and_blank_materials(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    session = write_log(env, date="2026-07-24", subject="Science", slug="plot")

    result = mod.cmd_append(append_args(env, session))
    assert result["written"] is True

    wb = openpyxl.load_workbook(env["workbook"])
    row = [c.value for c in wb["Science"][2]]
    assert row[:5] == ["7/24/2026", "9:15 AM", "10:05 AM", "Lesson 12: Mapping a Plot", "Brad"]
    # Materials and ISBN stay untouched so the dashboard's curricula table
    # is not filled with per-row repeats.
    assert row[5] is None and row[6] is None


def test_append_writes_a_hidden_stable_session_id(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    session = write_log(env, date="2026-07-24", subject="Science", slug="plot")

    result = mod.cmd_append(append_args(env, session))

    wb = openpyxl.load_workbook(env["workbook"])
    headers = [c.value for c in wb["Science"][1]]
    id_column = headers.index("Lesson Log ID") + 1
    assert result["session_id"] == "eliana/grade-3/2026/07/24/science-plot"
    assert wb["Science"].cell(row=2, column=id_column).value == result["session_id"]
    assert wb["Science"].column_dimensions[
        openpyxl.utils.get_column_letter(id_column)
    ].hidden is True


def test_correcting_a_start_time_does_not_create_a_duplicate(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    session = write_log(env, date="2026-07-24", subject="Science", slug="plot")
    mod.cmd_append(append_args(env, session))

    frontmatter = mod.read_frontmatter(session / "log.md")
    frontmatter["start_time"] = "9:20 AM"
    body = yaml.safe_dump(frontmatter, sort_keys=False).strip()
    (session / "log.md").write_text(
        f"---\n{body}\n---\n\n# Science\n", encoding="utf-8"
    )

    gaps = mod.cmd_gaps(gaps_args(env))
    assert gaps["pending"] == []
    assert gaps["already_logged"] == 1

    second = mod.cmd_append(append_args(env, session))
    assert second["written"] is False
    wb = openpyxl.load_workbook(env["workbook"])
    assert wb["Science"].max_row == 2


def test_two_sessions_can_share_a_date_and_start_time(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    first = write_log(env, date="2026-07-24", subject="Science", slug="plot")
    second = write_log(env, date="2026-07-24", subject="Science", slug="microscope")

    assert mod.cmd_append(append_args(env, first))["written"] is True
    assert mod.cmd_append(append_args(env, second))["written"] is True

    wb = openpyxl.load_workbook(env["workbook"])
    assert wb["Science"].max_row == 3


def test_append_places_notes_at_the_sheets_own_column(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    session = write_log(env, date="2026-07-24", subject="Science", slug="plot")

    result = mod.cmd_append(append_args(env, session, notes="Confused 6s and 9s."))
    assert result["notes_written"] is True

    wb = openpyxl.load_workbook(env["workbook"])
    headers = [c.value for c in wb["Science"][1]]
    row = [c.value for c in wb["Science"][2]]
    assert row[headers.index("Notes")] == "Confused 6s and 9s."


def test_append_skips_notes_when_the_sheet_has_no_such_column(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    session = write_log(env, date="2026-07-24", subject="Math", slug="math-45")

    result = mod.cmd_append(append_args(env, session, notes="Struggled with borrowing."))
    assert result["notes_written"] is False
    assert result["notes_skipped_no_column"] is True

    wb = openpyxl.load_workbook(env["workbook"])
    assert len([c.value for c in wb["Math"][2][:7] if c.value]) == 5


def test_append_is_a_no_op_when_the_row_already_exists(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    session = write_log(env, date="2026-07-24", subject="Science", slug="plot")

    first = mod.cmd_append(append_args(env, session))
    second = mod.cmd_append(append_args(env, session))

    assert first["written"] is True
    assert second["written"] is False
    wb = openpyxl.load_workbook(env["workbook"])
    assert wb["Science"].max_row == 2


def test_append_never_overwrites_a_hand_edited_description(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    session = write_log(env, date="2026-07-24", subject="Science", slug="plot")
    mod.cmd_append(append_args(env, session))

    wb = openpyxl.load_workbook(env["workbook"])
    wb["Science"].cell(row=2, column=4, value="My own wording")
    wb.save(env["workbook"])

    mod.cmd_append(append_args(env, session, description="Regenerated wording"))

    wb = openpyxl.load_workbook(env["workbook"])
    assert wb["Science"].cell(row=2, column=4).value == "My own wording"
    assert wb["Science"].max_row == 2


def test_append_refuses_a_log_with_no_times(tmp_path):
    mod = load_module()
    env = build_env(tmp_path)
    session = write_log(env, date="2026-07-24", subject="Science", slug="plot",
                        start=None, end=None)

    with pytest.raises(SystemExit) as excinfo:
        mod.cmd_append(append_args(env, session))
    assert "cannot be projected" in str(excinfo.value)


def test_hsd_without_time_is_a_configuration_error(tmp_path):
    mod = load_module()
    env = build_env(tmp_path, hsd=True, time=False)

    with pytest.raises(SystemExit) as excinfo:
        mod.cmd_gaps(gaps_args(env))
    assert "logging.hsd requires logging.time" in str(excinfo.value)


def test_hsd_without_lesson_logs_is_a_configuration_error(tmp_path):
    mod = load_module()
    env = build_env(tmp_path, hsd=True, lesson_log=False)

    with pytest.raises(SystemExit) as excinfo:
        mod.cmd_gaps(gaps_args(env))
    assert "logging.hsd requires logging.lesson_log" in str(excinfo.value)


def test_disabled_hsd_does_not_require_a_workbook(tmp_path):
    mod = load_module()
    env = build_env(tmp_path, hsd=False)
    Path(env["workbook"]).unlink()

    result = mod.cmd_gaps(gaps_args(env))
    assert result["hsd_enabled"] is False
    assert result["workbook"] is None
    assert result["pending"] == []


def test_build_row_lays_values_out_by_header_name():
    mod = load_module()
    headers = ["date", "start time", "end time", "description", "teacher",
               "materials", "isbn", "notes"]
    row = mod.build_row(headers, {"date": "7/24/2026", "description": "x", "notes": "y"})
    assert row == ["7/24/2026", "", "", "x", "", "", "", "y"]
